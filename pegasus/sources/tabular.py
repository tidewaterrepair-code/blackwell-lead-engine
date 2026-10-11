"""Generic CSV / Excel download (treasurer delinquent-tax lists, sales exports, violation logs...).

sources.yaml:
    vb-delinquent-taxes:
      type: csv
      kind: tax_delinquent
      url: https://example.gov/treasurer/delinquent.xlsx
      sheet: 0                       # Excel only (index or name)
      skip_rows: 0                   # rows above the header
      fields: {address: "Property Address", amount_due: "Amount Due", parcel_id: "GPIN"}
"""

from __future__ import annotations

import csv
import hashlib
import io
from datetime import datetime
from typing import Iterable

from .base import Adapter, LeadInput, RawItem, clean
from .permits_common import get, normalize_permit


def read_table(content: bytes, *, url: str = "", sheet=0, skip_rows: int = 0) -> list[dict]:
    if content[:4] == b"PK\x03\x04" or url.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.worksheets[sheet] if isinstance(sheet, int) else wb[sheet]
        rows = list(ws.iter_rows(values_only=True))[skip_rows:]
        if not rows:
            return []
        header = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(rows[0])]
        return [dict(zip(header, r)) for r in rows[1:] if any(v not in (None, "") for v in r)]
    text = content.decode("utf-8-sig", errors="replace")
    lines = text.splitlines()[skip_rows:]
    return [row for row in csv.DictReader(lines) if any((v or "").strip() for v in row.values())]


class TabularAdapter(Adapter):
    kind = "tax_delinquent"

    def fetch(self, since: datetime) -> Iterable[RawItem]:
        resp = self.client.get(self.config["url"])
        resp.raise_for_status()
        rows = read_table(resp.content, url=self.config["url"], sheet=self.config.get("sheet", 0),
                          skip_rows=int(self.config.get("skip_rows", 0)))
        for row in rows:
            row = {k: (v.isoformat() if isinstance(v, datetime) else v) for k, v in row.items()}
            ext = clean(get(row, self.config, "external_id")) or clean(get(row, self.config, "parcel_id"))
            if not ext:
                ext = hashlib.sha1(repr(sorted(row.items())).encode()).hexdigest()[:20]
            yield RawItem(external_id=f"{self.kind}:{ext}", payload=row)

    def normalize(self, payload: dict) -> LeadInput | None:
        item = normalize_permit(payload, {**self.config, "kind": self.kind}, default_city=self.default_city)
        if item and not item.external_id.startswith(self.kind):
            item.external_id = f"{self.kind}:{item.external_id}"
        return item
